import customtkinter as ctk
import sqlite3
import os
import traceback

from tkinter import messagebox
from tkinter import ttk

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from datetime import datetime

DB = "drivers.db"


class ReportsWindow(ctk.CTkFrame):

    def __init__(self, parent):
        super().__init__(parent)

        self.mode = "shifts"

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)

        # =====================================================
        # TITLE
        # =====================================================
        ctk.CTkLabel(
            self,
            text="📊 ERP ОТЧЁТНЫЙ ЦЕНТР",
            font=("Arial", 34, "bold"),
            text_color="white"
        ).grid(row=0, column=0, pady=15)

        # =====================================================
        # TOP PANEL
        # =====================================================
        top = ctk.CTkFrame(self, fg_color="#161616")
        top.grid(row=1, column=0, sticky="ew", padx=12, pady=8)

        ctk.CTkButton(
            top,
            text="🚍 СМЕНЫ",
            font=("Arial", 15, "bold"),
            height=40,
            command=self.load_shifts
        ).pack(side="left", padx=6, pady=8)

        ctk.CTkButton(
            top,
            text="🩺 МЕДКНИЖКИ",
            font=("Arial", 15, "bold"),
            height=40,
            command=self.load_med
        ).pack(side="left", padx=6, pady=8)

        ctk.CTkButton(
            top,
            text="📥 EXCEL",
            font=("Arial", 15, "bold"),
            height=40,
            fg_color="#2e8b57",
            hover_color="#3cb371",
            command=self.export_excel
        ).pack(side="left", padx=6, pady=8)

        # =====================================================
        # TABLE AREA
        # =====================================================
        table_frame = ctk.CTkFrame(self, fg_color="#0f0f0f")
        table_frame.grid(row=2, column=0, sticky="nsew", padx=12, pady=12)

        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        # =====================================================
        # STYLE
        # =====================================================
        style = ttk.Style()
        style.theme_use("default")

        style.configure(
            "Treeview",
            background="#1b1b1b",
            foreground="white",
            fieldbackground="#1b1b1b",
            borderwidth=0,
            rowheight=50,
            font=("Arial", 15)
        )

        style.configure(
            "Treeview.Heading",
            background="#1f4e79",
            foreground="white",
            font=("Arial", 16, "bold"),
            relief="flat"
        )

        style.map(
            "Treeview",
            background=[("selected", "#22577a")]
        )

        # =====================================================
        # TREEVIEW
        # =====================================================
        self.tree = ttk.Treeview(
            table_frame,
            show="headings"
        )

        self.tree.grid(row=0, column=0, sticky="nsew")

        scrollbar = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.tree.yview
        )

        self.tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.grid(row=0, column=1, sticky="ns")

        # =====================================================
        # TAG COLORS
        # =====================================================
        self.tree.tag_configure(
            "ok",
            background="#17351f",
            foreground="#7CFC98"
        )

        self.tree.tag_configure(
            "warn",
            background="#4a4316",
            foreground="#ffe066"
        )

        self.tree.tag_configure(
            "bad",
            background="#4a1616",
            foreground="#ff8a8a"
        )

        # =====================================================
        # COLUMNS
        # =====================================================
        self.tree["columns"] = ("c1", "c2", "c3", "c4", "c5")

        widths = [180, 320, 180, 180, 140]

        for c, w in zip(self.tree["columns"], widths):
            self.tree.column(c, width=w, anchor="center")

        self.load_shifts()

    # =====================================================
    # SAFE VALUE
    # =====================================================
    def safe(self, value):
        if value is None:
            return ""
        return str(value)
    
    def to_russian_date(self, date_str):
        try:
            return datetime.strptime(date_str, "%Y-%m-%d").strftime("%d.%m.%Y")
        except:
            return date_str

    # =====================================================
    # MED STATUS
    # =====================================================
    def med_status(self, date_str):

        try:
            med_date = datetime.strptime(
                date_str,
                "%Y-%m-%d"
            ).date()

            today = datetime.today().date()

            if med_date < today:
                return "❌ ПРОСРОЧЕНА", "bad"

            elif (med_date - today).days <= 30:
                return "⚠ СКОРО ЗАКОНЧИТСЯ", "warn"

            else:
                return "✔ АКТИВНА", "ok"

        except:
            return "❌ ОШИБКА", "bad"
        
    def to_russian_date(self, date_str):
        try:
            return datetime.strptime(date_str, "%Y-%m-%d").strftime("%d.%m.%Y")
        except:
            return date_str

    # =====================================================
    # LOAD SHIFTS
    # =====================================================
    def load_shifts(self):

        self.mode = "shifts"

        self.tree.delete(*self.tree.get_children())

        self.tree.heading("c1", text="Дата")
        self.tree.heading("c2", text="Водитель")
        self.tree.heading("c3", text="Автобус")
        self.tree.heading("c4", text="Маршрут")
        self.tree.heading("c5", text="Часы")

        conn = sqlite3.connect(DB)
        cur = conn.cursor()

        cur.execute("""
            SELECT
                s.shift_date,
                d.full_name,
                b.number,
                s.route,
                s.hours
            FROM shifts s
            LEFT JOIN drivers d ON d.id = s.driver_id
            LEFT JOIN buses b ON b.id = s.bus_id
            ORDER BY s.shift_date DESC
        """)

        rows = cur.fetchall()

        for row in rows:

            shift_date = self.to_russian_date(row[0])

            self.tree.insert(
                "",
                "end",
                values=(
                    shift_date,
                    self.safe(row[1]),
                    self.safe(row[2]),
                    self.safe(row[3]),
                    self.safe(row[4])
                )
            )

        conn.close()

    # =====================================================
    # LOAD MEDICAL BOOKS
    # =====================================================
    def load_med(self):

        self.mode = "med"

        self.tree.delete(*self.tree.get_children())

        self.tree.heading("c1", text="Водитель")
        self.tree.heading("c2", text="Медкнижка")
        self.tree.heading("c3", text="Срок действия")
        self.tree.heading("c4", text="Статус")
        self.tree.heading("c5", text="")

        conn = sqlite3.connect(DB)
        cur = conn.cursor()

        cur.execute("""
            SELECT
                full_name,
                med_book_number,
                med_expiry_date
            FROM drivers
            ORDER BY med_expiry_date ASC
        """)

        rows = cur.fetchall()

        for name, med, expiry in rows:

            status_text, tag = self.med_status(expiry)

            expiry = self.to_russian_date(expiry)

            self.tree.insert(
                "",
                "end",
                values=(
                    self.safe(name),
                    self.safe(med),
                    self.safe(expiry),
                    status_text,
                    ""
                ),
                tags=(tag,)
            )

        conn.close()

    # =====================================================
    # EXPORT EXCEL
    # =====================================================
    def export_excel(self):

        try:

            os.makedirs("reports", exist_ok=True)

            filename = (
                f"reports/ERP_REPORT_"
                f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )

            wb = Workbook()
            ws = wb.active

            ws.title = "ERP REPORT"

            # =================================================
            # COLORS
            # =================================================
            blue_fill = PatternFill(
                "solid",
                fgColor="1F4E79"
            )

            dark_fill = PatternFill(
                "solid",
                fgColor="D9EAF7"
            )

            red_fill = PatternFill(
                "solid",
                fgColor="F8CBCB"
            )

            yellow_fill = PatternFill(
                "solid",
                fgColor="FFF2CC"
            )

            green_fill = PatternFill(
                "solid",
                fgColor="D5F5E3"
            )

            # =================================================
            # FONTS
            # =================================================
            title_font = Font(
                size=20,
                bold=True
            )

            header_font = Font(
                size=13,
                bold=True,
                color="FFFFFF"
            )

            body_font = Font(
                size=12
            )

            center = Alignment(
                horizontal="center",
                vertical="center"
            )

            border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )

            # =================================================
            # TITLE
            # =================================================
            ws.merge_cells("A1:E1")

            ws["A1"] = "ERP ОТЧЁТ АВТОПАРКА"

            ws["A1"].font = title_font
            ws["A1"].alignment = center
            ws["A1"].fill = dark_fill

            ws.row_dimensions[1].height = 35

            ws.append([])

            # =================================================
            # EXPORT MED
            # =================================================
            if self.mode == "med":

                headers = [
                    "Водитель",
                    "Медкнижка",
                    "Срок",
                    "Статус"
                ]

                ws.append(headers)

                for row_id in self.tree.get_children():

                    vals = list(self.tree.item(row_id)["values"][:4])

                    ws.append(vals)

            # =================================================
            # EXPORT SHIFTS
            # =================================================
            else:

                headers = [
                    "Дата",
                    "Водитель",
                    "Автобус",
                    "Маршрут",
                    "Часы"
                ]

                ws.append(headers)

                for row_id in self.tree.get_children():

                    vals = self.tree.item(row_id)["values"]

                    ws.append(vals)

            # =================================================
            # HEADER STYLE
            # =================================================
            for cell in ws[3]:

                cell.font = header_font
                cell.fill = blue_fill
                cell.alignment = center
                cell.border = border

            # =================================================
            # BODY STYLE
            # =================================================
            for row in ws.iter_rows(min_row=4):

                for cell in row:

                    cell.font = body_font
                    cell.alignment = center
                    cell.border = border

            # =================================================
            # MED STATUS COLORS
            # =================================================
            if self.mode == "med":

                for row in range(4, ws.max_row + 1):

                    status = str(ws[f"D{row}"].value)

                    fill = green_fill

                    if "ПРОСРОЧЕНА" in status:
                        fill = red_fill

                    elif "СКОРО" in status:
                        fill = yellow_fill

                    for col in range(1, 5):
                        ws.cell(row=row, column=col).fill = fill

            # =================================================
            # AUTO WIDTH
            # =================================================
            for col in range(1, ws.max_column + 1):

                max_len = 0

                letter = get_column_letter(col)

                for row in range(1, ws.max_row + 1):

                    value = ws.cell(row=row, column=col).value

                    if value:
                        max_len = max(
                            max_len,
                            len(str(value))
                        )

                ws.column_dimensions[letter].width = max_len + 8

            # =================================================
            # FREEZE HEADER
            # =================================================
            ws.freeze_panes = "A4"

            # =================================================
            # SAVE
            # =================================================
            wb.save(filename)

            messagebox.showinfo(
                "Excel",
                f"✅ ERP отчёт создан\n\n{filename}"
            )

        except Exception as e:

            traceback.print_exc()

            messagebox.showerror(
                "Excel Error",
                str(e)
            )